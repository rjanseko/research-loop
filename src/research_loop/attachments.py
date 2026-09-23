from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import mimetypes
import re
from collections import Counter
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Any, Iterable

from pydantic import BaseModel, Field


class AttachmentMode(StrEnum):
    """How local attachments are exposed to agents.

    NORMALIZED is the default for comparative evals: every model sees the same
    deterministic extraction through local tools. MULTIMODAL additionally sends
    image bytes directly to models that support image input.
    """

    NORMALIZED = "normalized"
    MULTIMODAL = "multimodal"


class AttachmentKind(StrEnum):
    TEXT = "text"
    HTML = "html"
    PDF = "pdf"
    DOCX = "docx"
    CSV = "csv"
    XLSX = "xlsx"
    IMAGE = "image"
    UNKNOWN = "unknown"


class AttachmentLimits(BaseModel):
    max_file_bytes: int = 25 * 1024 * 1024
    max_text_chars: int = 500_000
    chunk_chars: int = 8_000
    max_chunks_per_attachment: int = 128
    max_csv_rows: int = 5_000
    max_sheet_rows: int = 2_000
    max_sheets: int = 32


class AttachmentChunk(BaseModel):
    chunk_id: str
    attachment_id: str
    ordinal: int
    locator: str
    kind: str = "text"
    text: str
    sha256: str


class AttachmentRecord(BaseModel):
    attachment_id: str
    name: str
    kind: AttachmentKind
    media_type: str
    sha256: str
    size_bytes: int
    extractor: str
    chunks: list[AttachmentChunk] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)
    truncated: bool = False
    requires_multimodal: bool = False
    extraction_error: str | None = None

    def manifest(self) -> dict[str, Any]:
        return {
            "attachment_id": self.attachment_id,
            "name": self.name,
            "kind": self.kind.value,
            "media_type": self.media_type,
            "sha256": self.sha256,
            "size_bytes": self.size_bytes,
            "extractor": self.extractor,
            "chunk_count": len(self.chunks),
            "metadata": self.metadata,
            "truncated": self.truncated,
            "requires_multimodal": self.requires_multimodal,
            "extraction_error": self.extraction_error,
        }


class AttachmentSearchHit(BaseModel):
    attachment_id: str
    name: str
    chunk_id: str
    locator: str
    score: float
    excerpt: str


class AttachmentIngestionError(RuntimeError):
    pass


@dataclass(frozen=True)
class _RawChunk:
    locator: str
    text: str
    kind: str = "text"


_TOKEN_RE = re.compile(r"[A-Za-z0-9_][A-Za-z0-9_\-]{1,}")
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
_TEXT_SUFFIXES = {".txt", ".md", ".rst", ".log", ".yaml", ".yml", ".toml", ".json", ".xml"}


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _clean_text(value: str) -> str:
    value = value.replace("\x00", "")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


def _chunk_text(text: str, *, max_chars: int, locator_prefix: str) -> list[_RawChunk]:
    text = _clean_text(text)
    if not text:
        return []
    chunks: list[_RawChunk] = []
    start = 0
    part = 1
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n\n", start, end), text.rfind("\n", start, end))
            if boundary > start + max_chars // 2:
                end = boundary
        piece = text[start:end].strip()
        if piece:
            locator = locator_prefix if len(text) <= max_chars else f"{locator_prefix}, part {part}"
            chunks.append(_RawChunk(locator=locator, text=piece))
            part += 1
        start = max(end, start + 1)
    return chunks


def _detect_kind(path: Path) -> AttachmentKind:
    suffix = path.suffix.casefold()
    if suffix in _TEXT_SUFFIXES:
        return AttachmentKind.TEXT
    if suffix in {".html", ".htm"}:
        return AttachmentKind.HTML
    if suffix == ".pdf":
        return AttachmentKind.PDF
    if suffix == ".docx":
        return AttachmentKind.DOCX
    if suffix == ".csv":
        return AttachmentKind.CSV
    if suffix in {".xlsx", ".xlsm"}:
        return AttachmentKind.XLSX
    if suffix in _IMAGE_SUFFIXES:
        return AttachmentKind.IMAGE
    return AttachmentKind.UNKNOWN


def _decode_text(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def _extract_text_file(path: Path, data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    text, encoding = _decode_text(data)
    if path.suffix.casefold() == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
    truncated = len(text) > limits.max_text_chars
    text = text[: limits.max_text_chars]
    return _chunk_text(text, max_chars=limits.chunk_chars, locator_prefix="document"), {"encoding": encoding}, truncated


def _extract_html(data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    text, encoding = _decode_text(data)
    title = None
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        if soup.title and soup.title.string:
            title = soup.title.string.strip()
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        extracted = soup.get_text("\n", strip=True)
        extractor = "beautifulsoup4"
    except ImportError:
        from html.parser import HTMLParser

        class _TextParser(HTMLParser):
            def __init__(self) -> None:
                super().__init__()
                self.parts: list[str] = []

            def handle_data(self, raw: str) -> None:
                if raw.strip():
                    self.parts.append(raw.strip())

        parser = _TextParser()
        parser.feed(text)
        extracted = "\n".join(parser.parts)
        extractor = "html.parser"
    truncated = len(extracted) > limits.max_text_chars
    extracted = extracted[: limits.max_text_chars]
    return _chunk_text(extracted, max_chars=limits.chunk_chars, locator_prefix="document"), {"encoding": encoding, "title": title, "html_extractor": extractor}, truncated


def _extract_pdf(data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise AttachmentIngestionError("PDF extraction requires the `attachments` extra (pypdf)") from exc

    reader = PdfReader(io.BytesIO(data))
    chunks: list[_RawChunk] = []
    chars = 0
    truncated = False
    pages_with_text = 0
    for page_index, page in enumerate(reader.pages, start=1):
        if chars >= limits.max_text_chars or len(chunks) >= limits.max_chunks_per_attachment:
            truncated = True
            break
        text = _clean_text(page.extract_text() or "")
        if text:
            pages_with_text += 1
        remaining = max(limits.max_text_chars - chars, 0)
        text = text[:remaining]
        page_chunks = _chunk_text(text, max_chars=limits.chunk_chars, locator_prefix=f"page {page_index}")
        for chunk in page_chunks:
            if len(chunks) >= limits.max_chunks_per_attachment:
                truncated = True
                break
            chunks.append(chunk)
        chars += len(text)
    metadata = {
        "page_count": len(reader.pages),
        "pages_with_text": pages_with_text,
        "text_page_ratio": (pages_with_text / len(reader.pages)) if reader.pages else 0.0,
    }
    if reader.metadata:
        for key in ("title", "author", "subject", "creator"):
            value = getattr(reader.metadata, key, None)
            if value:
                metadata[key] = str(value)
    return chunks, metadata, truncated


def _extract_docx(data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise AttachmentIngestionError("DOCX extraction requires the `attachments` extra (python-docx)") from exc

    doc = Document(io.BytesIO(data))
    sections: list[str] = []
    for p in doc.paragraphs:
        value = p.text.strip()
        if value:
            sections.append(value)
    for table_index, table in enumerate(doc.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            sections.append(f"[Table {table_index}]\n" + "\n".join(rows))
    text = "\n\n".join(sections)
    truncated = len(text) > limits.max_text_chars
    text = text[: limits.max_text_chars]
    props = doc.core_properties
    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "title": props.title or None,
        "author": props.author or None,
    }
    return _chunk_text(text, max_chars=limits.chunk_chars, locator_prefix="document"), metadata, truncated


def _rows_to_text(headers: list[str], rows: Iterable[list[Any]]) -> str:
    lines: list[str] = []
    if headers:
        lines.append(" | ".join(headers))
    for row in rows:
        lines.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def _extract_csv(data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    text, encoding = _decode_text(data)
    stream = io.StringIO(text)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(stream, dialect)
    rows = list(reader)
    if not rows:
        return [], {"encoding": encoding, "row_count": 0}, False
    headers = [str(x) for x in rows[0]]
    data_rows = rows[1:]
    truncated = len(data_rows) > limits.max_csv_rows
    data_rows = data_rows[: limits.max_csv_rows]
    chunks: list[_RawChunk] = []
    batch = 40
    for start in range(0, len(data_rows), batch):
        if len(chunks) >= limits.max_chunks_per_attachment:
            truncated = True
            break
        subset = data_rows[start : start + batch]
        body = _rows_to_text(headers, subset)
        chunks.append(_RawChunk(locator=f"rows {start + 2}-{start + 1 + len(subset)}", text=body, kind="table"))
    return chunks, {"encoding": encoding, "row_count": len(rows) - 1, "column_count": len(headers), "headers": headers[:100]}, truncated


def _extract_xlsx(data: bytes, limits: AttachmentLimits) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise AttachmentIngestionError("XLSX extraction requires the `attachments` extra (openpyxl)") from exc

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[_RawChunk] = []
    truncated = len(wb.sheetnames) > limits.max_sheets
    sheet_names = wb.sheetnames[: limits.max_sheets]
    row_counts: dict[str, int] = {}
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            first = next(rows_iter)
        except StopIteration:
            row_counts[sheet_name] = 0
            continue
        headers = ["" if x is None else str(x) for x in first]
        batch_rows: list[list[Any]] = []
        total = 0
        chunk_start = 2
        for row in rows_iter:
            total += 1
            if total > limits.max_sheet_rows:
                truncated = True
                break
            batch_rows.append(list(row))
            if len(batch_rows) >= 40:
                if len(chunks) >= limits.max_chunks_per_attachment:
                    truncated = True
                    break
                chunks.append(
                    _RawChunk(
                        locator=f"sheet {sheet_name!r}, rows {chunk_start}-{chunk_start + len(batch_rows) - 1}",
                        text=_rows_to_text(headers, batch_rows),
                        kind="table",
                    )
                )
                chunk_start += len(batch_rows)
                batch_rows = []
        if batch_rows and len(chunks) < limits.max_chunks_per_attachment:
            chunks.append(
                _RawChunk(
                    locator=f"sheet {sheet_name!r}, rows {chunk_start}-{chunk_start + len(batch_rows) - 1}",
                    text=_rows_to_text(headers, batch_rows),
                    kind="table",
                )
            )
        row_counts[sheet_name] = min(total, limits.max_sheet_rows)
        if len(chunks) >= limits.max_chunks_per_attachment:
            truncated = True
            break
    wb.close()
    return chunks, {"sheet_names": sheet_names, "row_counts": row_counts}, truncated


def _extract_image(data: bytes) -> tuple[list[_RawChunk], dict[str, Any], bool]:
    try:
        from PIL import Image
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise AttachmentIngestionError("Image metadata extraction requires the `attachments` extra (Pillow)") from exc

    with Image.open(io.BytesIO(data)) as image:
        metadata = {
            "width": image.width,
            "height": image.height,
            "format": image.format,
            "mode": image.mode,
            "frame_count": getattr(image, "n_frames", 1),
        }
    text = (
        "Image attachment. Pixel content is not converted to text in normalized mode. "
        f"Dimensions: {metadata['width']}x{metadata['height']}; format: {metadata['format']}."
    )
    return [_RawChunk(locator="entire image", text=text, kind="metadata")], metadata, False


class AttachmentCorpus:
    """Immutable-ish local corpus used by attachment tools during one research run."""

    def __init__(self, records: list[AttachmentRecord], paths: dict[str, Path]) -> None:
        self.records = records
        self._by_id = {record.attachment_id: record for record in records}
        self._paths = paths
        self._chunks = [chunk for record in records for chunk in record.chunks]
        self._doc_freq: Counter[str] = Counter()
        for chunk in self._chunks:
            for token in set(self._tokens(chunk.text)):
                self._doc_freq[token] += 1

    @classmethod
    def from_paths(
        cls,
        paths: Iterable[str | Path],
        *,
        limits: AttachmentLimits | None = None,
        strict: bool = True,
    ) -> "AttachmentCorpus":
        limits = limits or AttachmentLimits()
        records: list[AttachmentRecord] = []
        path_map: dict[str, Path] = {}
        for index, raw_path in enumerate(paths, start=1):
            path = Path(raw_path).expanduser().resolve()
            if not path.exists() or not path.is_file():
                message = f"attachment does not exist or is not a file: {path}"
                if strict:
                    raise AttachmentIngestionError(message)
                continue
            size = path.stat().st_size
            if size > limits.max_file_bytes:
                message = f"attachment exceeds max_file_bytes ({size} > {limits.max_file_bytes}): {path.name}"
                if strict:
                    raise AttachmentIngestionError(message)
                continue
            data = path.read_bytes()
            digest = _sha256_bytes(data)
            attachment_id = f"att-{index}-{digest[:12]}"
            kind = _detect_kind(path)
            media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            extraction_error: str | None = None
            extractor = kind.value
            raw_chunks: list[_RawChunk] = []
            metadata: dict[str, Any] = {}
            truncated = False
            requires_multimodal = False
            try:
                if kind is AttachmentKind.TEXT:
                    raw_chunks, metadata, truncated = _extract_text_file(path, data, limits)
                elif kind is AttachmentKind.HTML:
                    raw_chunks, metadata, truncated = _extract_html(data, limits)
                elif kind is AttachmentKind.PDF:
                    raw_chunks, metadata, truncated = _extract_pdf(data, limits)
                    requires_multimodal = bool(metadata.get("page_count")) and not bool(raw_chunks)
                    if metadata.get("text_page_ratio", 1.0) < 0.25:
                        requires_multimodal = True
                elif kind is AttachmentKind.DOCX:
                    raw_chunks, metadata, truncated = _extract_docx(data, limits)
                elif kind is AttachmentKind.CSV:
                    raw_chunks, metadata, truncated = _extract_csv(data, limits)
                elif kind is AttachmentKind.XLSX:
                    raw_chunks, metadata, truncated = _extract_xlsx(data, limits)
                elif kind is AttachmentKind.IMAGE:
                    raw_chunks, metadata, truncated = _extract_image(data)
                    requires_multimodal = True
                else:
                    extractor = "metadata-only"
                    requires_multimodal = True
                    raw_chunks = [
                        _RawChunk(
                            locator="file metadata",
                            text=(
                                "Unsupported binary attachment. Content was not decoded. "
                                f"Name: {path.name}; media type: {media_type}; size: {size} bytes."
                            ),
                            kind="metadata",
                        )
                    ]
            except Exception as exc:
                if strict:
                    raise
                extraction_error = f"{type(exc).__name__}: {exc}"
                raw_chunks = []
                requires_multimodal = True

            chunks: list[AttachmentChunk] = []
            total_chars = 0
            for ordinal, raw in enumerate(raw_chunks[: limits.max_chunks_per_attachment]):
                if total_chars >= limits.max_text_chars:
                    truncated = True
                    break
                text = raw.text[: max(limits.max_text_chars - total_chars, 0)]
                if not text:
                    continue
                chunks.append(
                    AttachmentChunk(
                        chunk_id=f"{attachment_id}:c{ordinal + 1}",
                        attachment_id=attachment_id,
                        ordinal=ordinal,
                        locator=raw.locator,
                        kind=raw.kind,
                        text=text,
                        sha256=_sha256_text(text),
                    )
                )
                total_chars += len(text)
            record = AttachmentRecord(
                attachment_id=attachment_id,
                name=path.name,
                kind=kind,
                media_type=media_type,
                sha256=digest,
                size_bytes=size,
                extractor=extractor,
                chunks=chunks,
                metadata=metadata,
                truncated=truncated,
                requires_multimodal=requires_multimodal,
                extraction_error=extraction_error,
            )
            records.append(record)
            path_map[attachment_id] = path
        return cls(records, path_map)

    @staticmethod
    def _tokens(text: str) -> list[str]:
        return [match.group(0).casefold() for match in _TOKEN_RE.finditer(text)]

    def manifest(self) -> list[dict[str, Any]]:
        return [record.manifest() for record in self.records]

    def prompt_manifest(self) -> list[dict[str, Any]]:
        return [
            {
                "attachment_id": r.attachment_id,
                "name": r.name,
                "kind": r.kind.value,
                "media_type": r.media_type,
                "sha256": r.sha256,
                "chunk_count": len(r.chunks),
                "requires_multimodal": r.requires_multimodal,
                "metadata": r.metadata,
                "truncated": r.truncated,
            }
            for r in self.records
        ]

    def list_attachments(self) -> list[dict[str, Any]]:
        """Return attachment metadata without leaking host filesystem paths."""
        return self.prompt_manifest()

    def get(self, attachment_id: str) -> AttachmentRecord:
        try:
            return self._by_id[attachment_id]
        except KeyError as exc:
            raise ValueError(f"unknown attachment_id {attachment_id!r}") from exc

    def read(self, attachment_id: str, *, start_chunk: int = 0, count: int = 5) -> dict[str, Any]:
        record = self.get(attachment_id)
        count = max(1, min(count, 10))
        start_chunk = max(start_chunk, 0)
        chunks = record.chunks[start_chunk : start_chunk + count]
        return {
            "attachment": record.manifest(),
            "start_chunk": start_chunk,
            "returned": len(chunks),
            "chunks": [
                {
                    "chunk_id": c.chunk_id,
                    "locator": c.locator,
                    "kind": c.kind,
                    "text": c.text,
                    "sha256": c.sha256,
                }
                for c in chunks
            ],
        }

    def search(
        self,
        query: str,
        *,
        attachment_id: str | None = None,
        top_k: int = 8,
    ) -> list[AttachmentSearchHit]:
        q_tokens = self._tokens(query)
        if not q_tokens:
            return []
        q_counts = Counter(q_tokens)
        candidates = self._chunks
        if attachment_id:
            self.get(attachment_id)
            candidates = [c for c in candidates if c.attachment_id == attachment_id]
        total_docs = max(len(self._chunks), 1)
        hits: list[AttachmentSearchHit] = []
        norm_query = query.casefold().strip()
        for chunk in candidates:
            tokens = self._tokens(chunk.text)
            if not tokens:
                continue
            counts = Counter(tokens)
            score = 0.0
            for token, q_weight in q_counts.items():
                tf = counts[token] / max(len(tokens), 1)
                idf = math.log((total_docs + 1) / (self._doc_freq[token] + 1)) + 1.0
                score += q_weight * tf * idf * 100.0
            if norm_query and norm_query in chunk.text.casefold():
                score += 5.0
            if score <= 0:
                continue
            record = self._by_id[chunk.attachment_id]
            excerpt = chunk.text[:1_500]
            hits.append(
                AttachmentSearchHit(
                    attachment_id=chunk.attachment_id,
                    name=record.name,
                    chunk_id=chunk.chunk_id,
                    locator=chunk.locator,
                    score=round(score, 6),
                    excerpt=excerpt,
                )
            )
        hits.sort(key=lambda h: h.score, reverse=True)
        return hits[: max(1, min(top_k, 20))]

    def multimodal_binary_specs(self) -> list[tuple[str, str, Path]]:
        """Return local binaries needed for the explicit multimodal benchmark lane.

        Images are always included. PDFs are included only when deterministic text
        extraction indicates that visual/document understanding is required (for example
        scanned PDFs). Text-bearing PDFs stay in the normalized attachment-tool layer.
        """
        result: list[tuple[str, str, Path]] = []
        for record in self.records:
            if record.kind is AttachmentKind.IMAGE or (
                record.kind is AttachmentKind.PDF and record.requires_multimodal
            ):
                result.append((record.attachment_id, record.media_type, self._paths[record.attachment_id]))
        return result


def build_attachment_toolset(corpus: AttachmentCorpus):
    """Build a run-local PydanticAI toolset over a normalized attachment corpus."""
    from pydantic_ai import FunctionToolset

    def list_attachments() -> list[dict[str, Any]]:
        """List local attachments available for this research run, including stable IDs and hashes."""
        return corpus.list_attachments()

    def search_attachments(
        query: str,
        attachment_id: str | None = None,
        top_k: int = 8,
    ) -> list[dict[str, Any]]:
        """Search extracted attachment text/tables. Use attachment_id to restrict to one file."""
        return [hit.model_dump(mode="json") for hit in corpus.search(query, attachment_id=attachment_id, top_k=top_k)]

    def read_attachment(
        attachment_id: str,
        start_chunk: int = 0,
        count: int = 5,
    ) -> dict[str, Any]:
        """Read normalized chunks from a local attachment by stable attachment ID."""
        return corpus.read(attachment_id, start_chunk=start_chunk, count=count)

    return FunctionToolset(tools=[list_attachments, search_attachments, read_attachment])


def build_multimodal_prompt(prompt: str, corpus: AttachmentCorpus) -> list[Any]:
    """Add local visual binaries as BinaryContent for the explicit multimodal lane.

    Text-bearing documents remain normalized through tools. Scanned/visual PDFs may be
    included only when extraction marks them as requiring multimodal understanding.
    """
    from pydantic_ai import BinaryContent

    parts: list[Any] = [prompt]
    for attachment_id, media_type, path in corpus.multimodal_binary_specs():
        # The model must see the bytes whose hash was recorded, not a file changed since.
        data = path.read_bytes()
        if _sha256_bytes(data) != corpus.get(attachment_id).sha256:
            raise AttachmentIngestionError(f"attachment {attachment_id} changed since ingestion; its recorded hash no longer matches")
        label = "image" if media_type.startswith("image/") else "document"
        parts.append(
            f"The following {label} corresponds to attachment_id={attachment_id}. "
            f"When using it as evidence, cite that attachment_id and an accurate locator."
        )
        parts.append(BinaryContent(data=data, media_type=media_type))
    return parts
