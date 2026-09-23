from __future__ import annotations

import csv
from pathlib import Path

from PIL import Image
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from research_loop.attachments import AttachmentCorpus, AttachmentKind
from research_loop.benchmarks.models import BenchmarkCaseSpec
from research_loop.schemas import SourceRef


def test_text_attachment_search_is_deterministic(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text(
        "Alpha project uses PostgreSQL for durable history.\n\n"
        "Beta project uses SQLite for local cache.",
        encoding="utf-8",
    )
    first = AttachmentCorpus.from_paths([path])
    second = AttachmentCorpus.from_paths([path])
    assert first.records[0].sha256 == second.records[0].sha256
    hits = first.search("PostgreSQL durable history")
    assert hits
    assert "PostgreSQL" in hits[0].excerpt
    assert hits[0].attachment_id == first.records[0].attachment_id


def test_structured_document_extractors(tmp_path: Path) -> None:
    html = tmp_path / "page.html"
    html.write_text("<html><head><title>T</title></head><body><h1>Heading</h1><p>HTML fact</p></body></html>")

    csv_path = tmp_path / "table.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["name", "score"])
        writer.writerow(["alice", 7])
        writer.writerow(["bob", 9])

    xlsx = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["name", "score"])
    ws.append(["alice", 7])
    wb.save(xlsx)

    docx_path = tmp_path / "brief.docx"
    doc = Document()
    doc.add_heading("Research Brief", level=1)
    doc.add_paragraph("DOCX fact")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "key"
    table.rows[0].cells[1].text = "value"
    doc.save(docx_path)

    pdf = tmp_path / "paper.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "PDF fact on page one")
    c.showPage()
    c.drawString(72, 720, "PDF fact on page two")
    c.save()

    corpus = AttachmentCorpus.from_paths([html, csv_path, xlsx, docx_path, pdf])
    kinds = {record.kind for record in corpus.records}
    assert {AttachmentKind.HTML, AttachmentKind.CSV, AttachmentKind.XLSX, AttachmentKind.DOCX, AttachmentKind.PDF} <= kinds
    assert corpus.search("HTML fact")
    assert corpus.search("bob 9")
    assert corpus.search("DOCX fact")
    pdf_hits = corpus.search("page two")
    assert pdf_hits and "page 2" in pdf_hits[0].locator


def test_image_is_metadata_only_in_normalized_corpus(tmp_path: Path) -> None:
    path = tmp_path / "chart.png"
    Image.new("RGB", (64, 32)).save(path)
    corpus = AttachmentCorpus.from_paths([path])
    record = corpus.records[0]
    assert record.kind is AttachmentKind.IMAGE
    assert record.requires_multimodal is True
    assert record.metadata["width"] == 64
    assert "Pixel content is not converted to text" in record.chunks[0].text


def test_benchmark_prompt_does_not_expose_host_path(tmp_path: Path) -> None:
    secret_dir = tmp_path / "private-host-dir"
    secret_dir.mkdir()
    attachment = secret_dir / "evidence.pdf"
    attachment.write_bytes(b"stub")
    case = BenchmarkCaseSpec(
        benchmark_id="gaia",
        case_id="x",
        objective="Answer from the file",
        attachments=[str(attachment)],
    )
    rendered = case.render_objective()
    assert "evidence.pdf" in rendered
    assert "private-host-dir" not in rendered
    assert str(secret_dir) not in rendered


def test_attachment_source_ref_requires_no_fake_url() -> None:
    source = SourceRef(
        attachment_id="att-1-deadbeef",
        locator="page 3",
        title="paper.pdf",
        source_type="attachment",
    )
    assert source.url is None
    assert source.is_attachment is True


def test_multimodal_input_refuses_a_file_changed_since_ingestion(tmp_path: Path) -> None:
    import pytest

    from research_loop.attachments import AttachmentIngestionError, build_multimodal_prompt

    path = tmp_path / "chart.png"
    Image.new("RGB", (8, 8), "red").save(path)
    corpus = AttachmentCorpus.from_paths([path])
    assert len(build_multimodal_prompt("prompt", corpus)) == 3  # prompt, label, image bytes
    Image.new("RGB", (8, 8), "blue").save(path)
    with pytest.raises(AttachmentIngestionError, match="changed since"):
        build_multimodal_prompt("prompt", corpus)


def test_attachment_ingestion_does_not_block_other_tasks(tmp_path: Path, monkeypatch) -> None:
    import asyncio
    import time
    from uuid import uuid4

    from research_loop.async_orchestrator import AsyncResearchLoop
    from research_loop.policy import ModelPolicy, ModelRoute
    from research_loop.schemas import ResearchConstraints, ResearchRole

    def slow_ingestion(*_args, **_kwargs):
        time.sleep(0.3)  # a large PDF or workbook
        return AttachmentCorpus([], {})

    monkeypatch.setattr(AttachmentCorpus, "from_paths", staticmethod(slow_ingestion))
    route = ModelRoute("test", 1, 1, 1_000)
    loop = AsyncResearchLoop(ModelPolicy("p", {role: route for role in ResearchRole}))

    async def main() -> int:
        ticks = 0

        async def ticker() -> None:
            nonlocal ticks
            while True:
                ticks += 1
                await asyncio.sleep(0.01)

        running = asyncio.create_task(ticker())
        await loop._load_attachments(uuid4(), ResearchConstraints(attachment_paths=["report.pdf"]))
        running.cancel()
        return ticks

    assert asyncio.run(main()) > 5  # the ticker kept running while the attachment was parsed
